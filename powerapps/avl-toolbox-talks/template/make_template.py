"""Builds TBX-SignOff-Template.xlsx: AV logo, heading block with named cells, declaration, and the
sign-off table heading at the named cell TableStart. Set up to print / save as A4 landscape PDF.

Run:  python3 make_template.py   (needs openpyxl and Pillow)
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

HERE = Path(__file__).parent
LOGO = HERE.parent / "assets" / "av-logo-small.png"
OUT = HERE / "TBX-SignOff-Template.xlsx"

GREEN = "005E38"
DECLARATION = ("I have read and understood the information in this toolbox talk, I have had the "
               "opportunity to ask questions, and I will follow the controls described.")
HEADERS = ["No.", "Name", "AVL ID", "Team", "Signed when", "Date & time", "Signature"]
WIDTHS = [6, 28, 12, 14, 20, 18, 26]

wb = Workbook()
ws = wb.active
ws.title = "Sign-off"
S = "'Sign-off'"

for i, w in enumerate(WIDTHS):
    ws.column_dimensions[chr(65 + i)].width = w

logo = XLImage(str(LOGO))
logo.height, logo.width = 70, int(70 * logo.width / logo.height)
ws.add_image(logo, "A1")

ws["C1"] = "Toolbox Talk Sign-off"
ws["C1"].font = Font(size=18, bold=True, color=GREEN)
ws.merge_cells("C1:G1")
ws["C2"] = "Australian Vintage · Buronga Hill Winery"
ws["C2"].font = Font(size=10, italic=True, color="555555")
ws.merge_cells("C2:G2")

thin = Side(style="thin", color="BBBBBB")
fields = [  # (row, label, name)
    (4, "Topic", "Topic"),
    (5, "Date", "MeetingDate"),
    (6, "Presenter", "Presenter"),
    (7, "Location", "Location"),
    (8, "Team", "Team"),
    (9, "Record", "MeetingID"),
]
for row, label, name in fields:
    ws.cell(row=row, column=2, value=label + ":").font = Font(bold=True)
    ws.cell(row=row, column=2).alignment = Alignment(horizontal="right")
    ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=7)
    cell = ws.cell(row=row, column=3)
    cell.border = Border(bottom=thin)
    wb.defined_names[name] = DefinedName(name, attr_text=f"{S}!$C${row}")

ws["A11"] = "Declaration"
ws["A11"].font = Font(bold=True, color=GREEN)
ws.merge_cells("A12:G13")
ws["A12"] = DECLARATION
ws["A12"].alignment = Alignment(wrap_text=True, vertical="top")
ws["A12"].font = Font(italic=True)
wb.defined_names["Declaration"] = DefinedName("Declaration", attr_text=f"{S}!$A$12")

TOP = 15
for i, h in enumerate(HEADERS):
    c = ws.cell(row=TOP, column=i + 1, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=GREEN)
    c.alignment = Alignment(vertical="center")
wb.defined_names["TableStart"] = DefinedName("TableStart", attr_text=f"{S}!$A${TOP}")

ws.freeze_panes = f"A{TOP + 1}"
ws.print_title_rows = f"{TOP}:{TOP}"
ws.page_setup.orientation = "landscape"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_options.horizontalCentered = True
ws.page_margins.left = ws.page_margins.right = 0.4
ws.oddFooter.left.text = "Signed electronically in the AVL Toolbox app"
ws.oddFooter.left.size = 8
ws.oddFooter.right.text = "Page &P of &N"
ws.oddFooter.right.size = 8

wb.save(OUT)
print("wrote", OUT)
